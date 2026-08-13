"""Migração única dos cadastros hoje em planilha para o banco.

Lê `Regras <mes>.xlsx` (abas MAPAS, REGRAS, Pendentes, LISTAS) e
`DE PARA GERAL.xlsx` (abas VERBAS, CADASTRAIS) e carrega:
  - usuarios (stub, um por e-mail de analista encontrado — sem login habilitado)
  - atributos_segmentacao
  - de_para_modelos "GERAL" + de_para_verbas
  - campos_cadastrais_mapa
  - clientes (MAPAS -> status ATIVO, Pendentes -> status PENDENTE)
  - regras_segmentacao (aba REGRAS, casada com os clientes recém-criados)

Não migra a base DS (fato) — isso é o pipeline de ingestão mensal (Fase 2), não
esta migração única de cadastro.

Uso:
    python scripts/migrate_from_excel.py "Regras julho.xlsx" "DE PARA GERAL.xlsx"

Roda dentro de uma única transação: se algo falhar, nada é gravado.

Limitações conhecidas (ver docs/plano.md, seção "Riscos"):
  - Clientes cujo "Modelo De para" na aba MAPAS não é GERAL caem de volta para o
    modelo GERAL com um aviso — nos dados usados para desenhar este script não havia
    nenhum caso real de De/Para próprio para validar esse caminho.
  - Cada linha da aba Pendentes é importada como um cliente PENDENTE à parte (mesma
    granularidade da aba MAPAS); a coluna "Nome Ctt" (que agrupa várias linhas, ex:
    várias unidades do "MERCADO LIVRE") é preservada em `observacao`, não usada para
    agrupar automaticamente — confirmar com o faturista se o agrupamento estava certo.
"""

import argparse
import secrets
import sys
from collections import defaultdict

import openpyxl

sys.path.insert(0, ".")  # permite `python scripts/migrate_from_excel.py` a partir de backend/

from app.core.crypto import encrypt_secret
from app.core.database import SessionLocal
from app.core.security import hash_password
from app.models.atributo import AtributoSegmentacao
from app.models.cliente import Cliente
from app.models.de_para import CampoCadastralMapa, DeParaModelo, DeParaVerba
from app.models.enums import ComparadorCondicao, LogicaRegra, PapelUsuario, StatusCliente, StatusFolha
from app.models.regra import RegraCondicao, RegraSegmentacao
from app.models.usuario import Usuario


def s(value) -> str | None:
    """Normaliza uma célula: remove espaços nas pontas, transforma vazio em None."""
    if value is None:
        return None
    texto = str(value).strip()
    return texto or None


def parece_email(valor: str | None) -> bool:
    """Checagem simples (não é validação RFC completa) para filtrar lixo que às
    vezes aparece na coluna 'Analista Responsável' da planilha (ex: '0',
    'faturamento rj', notas soltas) em vez de um e-mail real."""
    return bool(valor) and "@" in valor and "." in valor.split("@", 1)[1]


def eh_sim(value) -> bool:
    texto = s(value)
    return bool(texto) and texto.strip().upper().startswith("SIM")


def eh_folha_aberta(value) -> bool:
    texto = s(value) or ""
    return "ABERTA" in texto.upper()


def para_int(value) -> int | None:
    texto = s(value)
    if texto is None:
        return None
    try:
        return int(float(texto))
    except ValueError:
        return None


def linhas(ws, min_row=2):
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if all(c is None for c in row):
            continue
        yield row


def carregar_usuarios(db, emails: set[str]) -> dict[str, int]:
    mapa = {}
    criados = 0
    for email in sorted(e for e in emails if e and parece_email(e)):
        email_norm = email.lower()
        usuario = db.query(Usuario).filter(Usuario.email == email_norm).first()
        if usuario is None:
            usuario = Usuario(
                nome=email_norm.split("@")[0].replace(".", " ").title(),
                email=email_norm,
                senha_hash=hash_password(secrets.token_urlsafe(16)),
                papel=PapelUsuario.ANALISTA,
                ativo=False,  # precisa de reset de senha antes do primeiro login
            )
            db.add(usuario)
            db.flush()
            criados += 1
        mapa[email_norm] = usuario.id
    print(f"  usuarios: {criados} stub(s) criado(s), {len(mapa)} distintos no total")
    return mapa


def carregar_atributos(db, codigos: set[str]) -> None:
    existentes = {a.codigo for a in db.query(AtributoSegmentacao).all()}
    criados = 0
    for codigo in sorted(c for c in codigos if c):
        codigo_norm = codigo.strip().upper()
        if codigo_norm in existentes:
            continue
        db.add(AtributoSegmentacao(codigo=codigo_norm))
        existentes.add(codigo_norm)
        criados += 1
    print(f"  atributos_segmentacao: {criados} criado(s), {len(existentes)} distintos no total")


def carregar_de_para_geral(db, wb_de_para) -> DeParaModelo:
    modelo = db.query(DeParaModelo).filter(DeParaModelo.nome == "GERAL").first()
    if modelo is None:
        modelo = DeParaModelo(nome="GERAL")
        db.add(modelo)
        db.flush()

    ws = wb_de_para["VERBAS"]
    total = 0
    for cod, descri, eventos, grupo, ordem_grupo, ordem_item in linhas(ws):
        db.add(
            DeParaVerba(
                de_para_modelo_id=modelo.id,
                verba_codigo=s(cod),
                descricao_original=s(descri),
                evento_exibicao=s(eventos) or s(descri),
                grupo=s(grupo) or "Sem Grupo",
                ordem_grupo=para_int(ordem_grupo) or 0,
                ordem_item=para_int(ordem_item) or 0,
            )
        )
        total += 1
    print(f"  de_para_verbas (GERAL): {total} item(ns)")
    return modelo


def carregar_campos_cadastrais(db, wb_de_para) -> None:
    ws = wb_de_para["CADASTRAIS"]
    total = 0
    for row in linhas(ws):
        campo, ordem = row[0], row[1]
        if s(campo) is None:
            continue
        db.add(CampoCadastralMapa(campo=s(campo), ordem=para_int(ordem) or 0))
        total += 1
    print(f"  campos_cadastrais_mapa: {total} campo(s)")


def resolver_de_para_modelo_id(valor, geral_modelo: DeParaModelo, cliente_nome: str, avisos: list[str]) -> tuple[int, str | None]:
    """Retorna (de_para_modelo_id, nota_para_observacao).

    Quando o cliente pede um De/Para próprio (valor != GERAL), este script ainda não
    sabe montar esse modelo customizado (não havia exemplo suficiente nos dados para
    validar o formato) — cai para GERAL, mas grava uma nota tanto no aviso final
    quanto na observação do próprio cliente, para não perder a informação.
    """
    texto = s(valor)
    if texto is None or texto.upper() == "GERAL":
        return geral_modelo.id, None
    avisos.append(f"Cliente '{cliente_nome}': pede De/Para próprio '{texto}' — usando GERAL, revisar manualmente")
    return geral_modelo.id, f"PENDENTE: precisa de De/Para próprio '{texto}' (migrado com GERAL provisoriamente)"


def carregar_clientes_mapas(
    db, ws, usuarios: dict[str, int], geral_modelo: DeParaModelo, avisos: list[str], segmentacao_mapa_legada: dict[int, str]
) -> dict[tuple[str, str], int]:
    """`segmentacao_mapa_legada` é preenchido aqui (cliente_id -> segmentação que a
    planilha tinha) para `carregar_regras` saber que atributo cada regra migrada
    deve testar — não existe mais coluna de segmentação no cadastro do cliente
    (ver models/regra.py: cada regra agora carrega seu próprio atributo)."""
    resolvidos: dict[tuple[str, str], int] = {}
    total = 0
    for row in linhas(ws):
        (
            _regra,
            negocio,
            cliente_nome,
            _seg_email,
            seg_mapa,
            de_para,
            modelo_mapa,
            analista_email,
            folha,
            aguardo_po,
            site,
            login,
            senha,
            observacao,
            _chave,
        ) = row[:15]

        negocio_n, nome_n = s(negocio), s(cliente_nome)
        if negocio_n is None or nome_n is None:
            continue

        de_para_modelo_id, nota_de_para = resolver_de_para_modelo_id(de_para, geral_modelo, nome_n, avisos)
        observacao_final = " | ".join(p for p in [s(observacao), nota_de_para] if p) or None

        cliente = Cliente(
            negocio=negocio_n,
            nome=nome_n,
            status=StatusCliente.ATIVO,
            de_para_modelo_id=de_para_modelo_id,
            modelo_mapa_codigo=(s(modelo_mapa) or "GERAL").upper(),
            analista_responsavel_id=usuarios.get((s(analista_email) or "").lower()) if s(analista_email) else None,
            folha=StatusFolha.ABERTA if eh_folha_aberta(folha) else StatusFolha.FECHADA,
            aguardo_po=eh_sim(aguardo_po),
            portal_site=s(site),
            portal_login_cifrado=encrypt_secret(s(login)),
            portal_senha_cifrada=encrypt_secret(s(senha)),
            observacao=observacao_final,
        )
        db.add(cliente)
        db.flush()
        resolvidos[(negocio_n, nome_n)] = cliente.id
        segmentacao_mapa_legada[cliente.id] = (s(seg_mapa) or "GERAL").upper()
        total += 1
    print(f"  clientes (MAPAS -> ATIVO): {total}")
    return resolvidos


def carregar_clientes_pendentes(
    db,
    ws,
    usuarios: dict[str, int],
    geral_modelo: DeParaModelo,
    avisos: list[str],
    resolvidos: dict[tuple[str, str], int],
    segmentacao_mapa_legada: dict[int, str],
) -> None:
    total = 0
    for row in linhas(ws):
        (
            negocio,
            cliente_nome,
            nome_ctt,
            tipo_segmentacao,
            de_para,
            modelo_mapa,
            analista_email,
            _data_envio,
            folha,
            _emissao_automatica,
            aguardo_po,
            site,
            login,
            senha,
            observacao,
        ) = row[:15]

        negocio_n, nome_n = s(negocio), s(cliente_nome)
        if negocio_n is None or nome_n is None:
            continue
        if (negocio_n, nome_n) in resolvidos:
            continue  # já existe como ativo, não duplica

        de_para_modelo_id, nota_de_para = resolver_de_para_modelo_id(de_para, geral_modelo, nome_n, avisos)
        partes_obs = [p for p in [s(observacao), nota_de_para,
                                   f"Grupo (Nome Ctt): {s(nome_ctt)}" if s(nome_ctt) else None,
                                   f"Tipo Segmentação (texto livre da planilha): {s(tipo_segmentacao)}" if s(tipo_segmentacao) else None] if p]

        cliente = Cliente(
            negocio=negocio_n,
            nome=nome_n,
            status=StatusCliente.PENDENTE,
            de_para_modelo_id=de_para_modelo_id,
            modelo_mapa_codigo=(s(modelo_mapa) or "GERAL").upper(),
            analista_responsavel_id=usuarios.get((s(analista_email) or "").lower()) if s(analista_email) else None,
            folha=StatusFolha.ABERTA if eh_folha_aberta(folha) else StatusFolha.FECHADA,
            aguardo_po=eh_sim(aguardo_po),
            portal_site=s(site),
            portal_login_cifrado=encrypt_secret(s(login)),
            portal_senha_cifrada=encrypt_secret(s(senha)),
            observacao=" | ".join(partes_obs) or None,
        )
        db.add(cliente)
        db.flush()
        resolvidos[(negocio_n, nome_n)] = cliente.id
        segmentacao_mapa_legada[cliente.id] = "GERAL"
        total += 1
    print(f"  clientes (Pendentes -> PENDENTE): {total}")


def carregar_regras(
    db,
    ws,
    usuarios: dict[str, int],
    clientes: dict[tuple[str, str], int],
    avisos: list[str],
    segmentacao_mapa_legada: dict[int, str],
) -> None:
    total = 0
    nao_casados = 0
    ordem_por_cliente: dict[int, int] = defaultdict(int)
    for row in linhas(ws):
        (
            _regra,
            negocio,
            cliente_nome,
            nome_exibicao,
            valor,
            email_responsavel,
            data_envio,
            envio_automatico,
            analista_email,
            _chave,
        ) = row[:10]

        negocio_n, nome_n = s(negocio), s(cliente_nome)
        if negocio_n is None or nome_n is None:
            continue

        cliente_id = clientes.get((negocio_n, nome_n))
        if cliente_id is None:
            # fallback: tenta achar ignorando maiúsculas/minúsculas
            alvo = next(
                (cid for (neg, nome), cid in clientes.items() if neg == negocio_n and nome.upper() == nome_n.upper()),
                None,
            )
            cliente_id = alvo

        if cliente_id is None:
            nao_casados += 1
            avisos.append(f"Regra sem cliente correspondente: negocio={negocio_n!r} cliente={nome_n!r}")
            continue

        ordem = ordem_por_cliente[cliente_id]
        ordem_por_cliente[cliente_id] += 1

        db.add(
            RegraSegmentacao(
                cliente_id=cliente_id,
                ordem=ordem,
                logica=LogicaRegra.E,
                nome_exibicao=s(nome_exibicao) or "",
                email_responsavel=s(email_responsavel) or "",
                dia_envio=para_int(data_envio),
                envio_automatico=eh_sim(envio_automatico),
                analista_id=usuarios.get((s(analista_email) or "").lower()) if s(analista_email) else None,
                # A planilha só tinha um valor por regra, e o atributo era um único
                # por cliente (coluna "Segmentação Mapa" da aba MAPAS) — então cada
                # regra migra com uma condição só, "atributo IGUAL valor". Condições
                # extras (E/OU) são cadastradas depois pela tela.
                condicoes=[
                    RegraCondicao(
                        atributo=segmentacao_mapa_legada.get(cliente_id, "GERAL"),
                        comparador=ComparadorCondicao.IGUAL,
                        valor=s(valor) or "",
                    )
                ],
            )
        )
        total += 1
    print(f"  regras_segmentacao: {total} carregada(s), {nao_casados} não casada(s) com nenhum cliente")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("regras_xlsx", help='Caminho para "Regras <mes>.xlsx"')
    parser.add_argument("de_para_xlsx", help='Caminho para "DE PARA GERAL.xlsx"')
    args = parser.parse_args()

    print(f"Lendo {args.regras_xlsx} ...")
    wb_regras = openpyxl.load_workbook(args.regras_xlsx, data_only=True)
    print(f"Lendo {args.de_para_xlsx} ...")
    wb_de_para = openpyxl.load_workbook(args.de_para_xlsx, data_only=True)

    ws_mapas = wb_regras["MAPAS"]
    ws_regras = wb_regras["REGRAS"]
    ws_pendentes = wb_regras["Pendentes"]
    ws_listas = wb_regras["LISTAS"]

    avisos: list[str] = []

    # --- coleta prévia para usuarios e atributos, antes de qualquer insert ---
    emails = set()
    emails_invalidos = set()
    for row, coluna in (
        *((r, 7) for r in linhas(ws_mapas)),  # Analista Responsável
        *((r, 6) for r in linhas(ws_pendentes)),  # Analista Responsável
        *((r, 8) for r in linhas(ws_regras)),  # Analista
    ):
        valor = s(row[coluna])
        if valor is None:
            continue
        if parece_email(valor):
            emails.add(valor)
        else:
            emails_invalidos.add(valor)
    if emails_invalidos:
        avisos.append(
            f"{len(emails_invalidos)} valor(es) na coluna Analista não parecem e-mail e foram ignorados: "
            f"{sorted(emails_invalidos)}"
        )

    codigos_atributos = {s(row[0]) for row in linhas(ws_listas)} | {"COLABORADOR", "GERAL"}
    for row in linhas(ws_mapas):
        codigos_atributos.add(s(row[3]))
        codigos_atributos.add(s(row[4]))
    for row in linhas(ws_pendentes):
        codigos_atributos.add(s(row[3]))

    db = SessionLocal()
    try:
        print("Carregando cadastros de apoio...")
        usuarios = carregar_usuarios(db, emails)
        carregar_atributos(db, codigos_atributos)

        print("Carregando De/Para GERAL...")
        geral_modelo = carregar_de_para_geral(db, wb_de_para)
        carregar_campos_cadastrais(db, wb_de_para)

        print("Carregando clientes...")
        segmentacao_mapa_legada: dict[int, str] = {}
        clientes = carregar_clientes_mapas(db, ws_mapas, usuarios, geral_modelo, avisos, segmentacao_mapa_legada)
        carregar_clientes_pendentes(db, ws_pendentes, usuarios, geral_modelo, avisos, clientes, segmentacao_mapa_legada)

        print("Carregando regras de segmentação...")
        carregar_regras(db, ws_regras, usuarios, clientes, avisos, segmentacao_mapa_legada)

        db.commit()
        print("\nMigração concluída e commitada.")
    except Exception:
        db.rollback()
        print("\nERRO — nada foi gravado (rollback completo).")
        raise
    finally:
        db.close()

    if avisos:
        print(f"\n{len(avisos)} aviso(s) para revisar:")
        for aviso in avisos:
            print(f"  - {aviso}")


if __name__ == "__main__":
    main()
